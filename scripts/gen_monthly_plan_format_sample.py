# -*- coding: utf-8 -*-
"""월간 업무계획 hwpx 양식 설계용 참고 샘플(docx) 생성기.

목적: 현재 hwpx 산출물이 표·정렬 없이 문단만 이어져 가독성이 낮다.
      사서가 새 hwpx 양식을 설계할 때 참고할 "이렇게 보였으면 좋겠다" 견본을 만든다.

산출:
  산출물/[양식샘플]2026년_10월_도서관_월간업무계획.docx      — 전체본
  산출물/[양식샘플]2026년_10월_월간업무계획_요약보고.docx    — 최고관리자용 1~2쪽 요약본
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(BASE, "산출물")

# ── 색상 팔레트 (공문서 톤 — 짙은 청록 + 무채색) ──────────────────
C_MAIN = "1E6E5C"      # 제목·강조선
C_HEAD_BG = "E8EFEC"   # 표 머리 배경
C_SUB_BG = "F4F6F5"    # 소계·합계 행 배경
C_BORDER = "9AA5A0"    # 표 테두리
C_TEXT = "20242A"
C_MUTED = "6B7078"

FONT = "맑은 고딕"


# ── 저수준 유틸 ────────────────────────────────────────────────
def set_run(run, size=10, bold=False, color=C_TEXT, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return run


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


def cell_borders(cell, color=C_BORDER, sz=6):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tcPr.append(borders)


def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + name)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def vcenter(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    e = OxmlElement("w:vAlign")
    e.set(qn("w:val"), "center")
    tcPr.append(e)


def para_border_left(p, color=C_MAIN, sz=18):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    e = OxmlElement("w:left")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "6")
    e.set(qn("w:color"), color)
    pbdr.append(e)
    pPr.append(pbdr)


def para_border_bottom(p, color=C_MAIN, sz=18):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "6")
    e.set(qn("w:color"), color)
    pbdr.append(e)
    pPr.append(pbdr)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


# ── 문서 요소 ──────────────────────────────────────────────────
def doc_title(doc, text, sub=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4 if sub else 10)
    set_run(p.add_run(text), size=19, bold=True, color="17322A")
    para_border_bottom(p, sz=20)
    if sub:
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(10)
        set_run(q.add_run(sub), size=9.5, color=C_MUTED)


def h1(doc, text, space_before=16):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(5)
    set_run(p.add_run(text), size=12.5, bold=True, color=C_MAIN)
    para_border_left(p)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run(text), size=10.5, bold=True, color="3A3F46")
    return p


def body(doc, text, size=10, indent=0.4, color=C_TEXT, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.35
    set_run(p.add_run(text), size=size, color=color, bold=bold)
    return p


def note(doc, text):
    return body(doc, text, size=8.5, indent=0.4, color=C_MUTED)


def caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(text), size=9, bold=True, color="3A3F46")
    return p


def make_table(doc, headers, rows, widths, aligns=None, font_size=9,
               head_size=9, emphasize_last_row=False, header_align=None):
    """헤더 1행 + 본문 n행 표. widths=cm 리스트, aligns=('l','c','r') 리스트."""
    aligns = aligns or ["l"] * len(headers)
    amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
            "r": WD_ALIGN_PARAGRAPH.RIGHT}
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False

    def fill(cell, text, w, align, bold, size, bg):
        cell.width = Cm(w)
        cell_borders(cell)
        cell_margins(cell)
        vcenter(cell)
        if bg:
            shade(cell, bg)
        p = cell.paragraphs[0]
        p.alignment = amap[align]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.25
        set_run(p.add_run(str(text)), size=size, bold=bold)

    hdr = t.rows[0]
    repeat_header(hdr)
    for i, htext in enumerate(headers):
        fill(hdr.cells[i], htext, widths[i], (header_align or "c"), True, head_size, C_HEAD_BG)

    for r_i, row in enumerate(rows):
        cells = t.add_row().cells
        last = emphasize_last_row and r_i == len(rows) - 1
        for i, val in enumerate(row):
            fill(cells[i], val, widths[i], aligns[i], last, font_size,
                 C_SUB_BG if last else None)
    return t


def kv_table(doc, pairs, label_w=2.6, value_w=6.0, cols=2, font_size=9):
    """문서 정보용 라벨/값 표. cols=2면 (라벨,값) 2쌍이 한 줄에 들어간다."""
    ncol = cols * 2
    widths = ([label_w, value_w] * cols)
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i in range(0, len(pairs), cols):
        chunk = list(pairs[i:i + cols])
        while len(chunk) < cols:
            chunk.append(("", ""))
        cells = t.add_row().cells
        for j, (k, v) in enumerate(chunk):
            lc, vc = cells[j * 2], cells[j * 2 + 1]
            for c, w in ((lc, widths[j * 2]), (vc, widths[j * 2 + 1])):
                c.width = Cm(w)
                cell_borders(c)
                cell_margins(c)
                vcenter(c)
            shade(lc, C_HEAD_BG)
            for c, text, bold, align in ((lc, k, True, WD_ALIGN_PARAGRAPH.CENTER),
                                         (vc, v, False, WD_ALIGN_PARAGRAPH.LEFT)):
                p = c.paragraphs[0]
                p.alignment = align
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                set_run(p.add_run(str(text)), size=font_size, bold=bold)
    return t


def kpi_row(doc, items):
    """요약본 상단 핵심 지표 타일 — 라벨 행 + 값 행 2단 표."""
    t = doc.add_table(rows=2, cols=len(items))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    w = round(17.0 / len(items), 2)
    for i, (label, value, unit) in enumerate(items):
        lc = t.rows[0].cells[i]
        vc = t.rows[1].cells[i]
        for c in (lc, vc):
            c.width = Cm(w)
            cell_borders(c)
            cell_margins(c, top=70, bottom=70)
            vcenter(c)
        shade(lc, C_HEAD_BG)
        p = lc.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        set_run(p.add_run(label), size=8.5, bold=True, color="3A3F46")
        q = vc.paragraphs[0]
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_before = Pt(2)
        q.paragraph_format.space_after = Pt(2)
        set_run(q.add_run(value), size=14, bold=True, color=C_MAIN)
        set_run(q.add_run((" " + unit) if unit else ""), size=8.5, color=C_MUTED)
    return t


def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(2)
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    return doc


def footer_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(text), size=8.5, color=C_MUTED)
    para_border_bottom(p, color="C9CEC4", sz=6)


# ══════════════════════════════════════════════════════════════
# 공통 데이터 (두 문서가 같은 수치를 쓴다)
# ══════════════════════════════════════════════════════════════
DOMAIN_SUMMARY = [
    ("DM-01 장서개발", "희망도서 주간 수서 4회 / 11월 정기수서 사전 후보 선정 / 장서현황 월간 보고",
     "3,600,000", "1건"),
    ("DM-02 이용자응대", "키오스크 응대 상시 / 상호대차 일 평균 12건 / 하반기 만족도 조사 실시",
     "0", "-"),
    ("DM-03 독서문화", "10월 독서진흥행사 6종 운영 / 문화가 있는 날(10. 28.) / 순회문고 3회차 배치",
     "1,430,000", "1건"),
    ("DM-04 평생학습", "하반기 강좌 3종 운영(9~11월) 6~9주차 / 강사비 월 지급 / 중간 만족도 점검",
     "1,800,000", "-"),
    ("DM-05 홍보협력", "소식지 4분기호 편집 / 운영위원회 2차 회의(10. 15.) / 공모사업 1건 신청",
     "400,000", "1건"),
    ("기획담당 직접", "9월 실적 통계 집계·보고 / 2027년 주요업무계획 초안 착수 / 성과지표 3분기 점검",
     "0", "-"),
]

BUDGET_ROWS = [
    ("독서문화프로그램운영", "3.도서관독서진흥행사", "[2100143]일반수용비", "4,740,000",
     "2,260,960", "830,000", "1,649,040"),
    ("독서문화프로그램운영", "3.도서관독서진흥행사", "[2100605]강사수당", "1,800,000",
     "1,300,000", "300,000", "200,000"),
    ("독서문화프로그램운영", "1.도서관특화프로그램운영", "[2100605]강사수당", "1,800,000",
     "900,000", "300,000", "600,000"),
    ("독서문화프로그램운영", "2.도서관협력및홍보", "[2100612]위원회수당", "400,000",
     "0", "400,000", "0"),
    ("평생학습프로그램운영", "1.평생학습강좌운영", "[2100605]강사수당", "9,000,000",
     "4,500,000", "1,800,000", "2,700,000"),
    ("종합자료실운영", "2.장서확충", "[2110101]도서구입비", "36,000,000",
     "21,400,000", "3,600,000", "11,000,000"),
    ("합  계", "", "", "53,740,000", "30,360,960", "7,230,000", "16,149,040"),
]

ESCALATIONS = [
    ("1", "DM-01 / B-01", "11월 정기수서 예산 잔액 부족 예상",
     "도서구입비 잔액 ₩11,000,000 대비 11월 정기수서 소요 예상액 ₩13,200,000. "
     "부족액 ₩2,200,000에 대해 동일 단위과제카드 내 인접 사업항목 유용 여부 결정 필요",
     "10. 10.(토)까지"),
    ("2", "DM-03 / D-02", "문화가 있는 날 외부 공연 섭외비 초과",
     "당초 계상 ₩300,000 대비 섭외 가능 공연팀 최저 견적 ₩450,000. "
     "초과분 ₩150,000 일반수용비 전용 또는 프로그램 축소 중 택일 필요",
     "10. 8.(목)까지"),
    ("3", "DM-05 / F-03", "경상남도교육청 독서문화 공모사업 신청 여부",
     "신청 마감 10. 20.(화), 사업비 ₩5,000,000 규모. 선정 시 11~12월 추가 행사 2종 "
     "운영 필요 — 인력 여건 검토 후 신청 여부 결정 요청",
     "10. 13.(화)까지"),
]


# ══════════════════════════════════════════════════════════════
# 문서 A — 전체본
# ══════════════════════════════════════════════════════════════
def build_full():
    doc = new_doc()

    doc_title(doc, "2026년 10월 도서관 월간 업무계획")
    kv_table(doc, [
        ("대상 기간", "2026. 10. 1.(목) ~ 10. 31.(토)"),
        ("작성 부서", "창녕도서관 기획업무팀"),
        ("작 성 자", "기획업무팀 기획담당"),
        ("작 성 일", "2026. 9. 25.(금)"),
        ("문서 상태", "초안 — 최고관리자 검토 전"),
        ("총 소요예산", "금7,230,000원"),
    ], label_w=2.4, value_w=6.1, cols=2)

    # 1. 총괄 요약
    h1(doc, "1. 총괄 요약", space_before=18)
    body(doc, "가. 10월은 「독서의 달」 후속 기간으로 독서진흥행사 6종과 하반기 평생학습 강좌 "
              "3종을 병행 운영합니다.")
    body(doc, "나. 11월 정기도서수서 사전 준비(후보 선정·복본조사)를 10월 중 완료하여 "
              "11월 1일 자료심의위원회 개최에 대비합니다.")
    body(doc, "다. 최고관리자 확인·승인이 필요한 항목은 3건이며, 세부 내용은 5절에 정리하였습니다.")

    caption(doc, "[표 1] 도메인별 계획 요약")
    make_table(
        doc,
        ["도메인", "핵심 업무", "소요예산(원)", "확인필요"],
        DOMAIN_SUMMARY,
        widths=[2.8, 9.4, 2.6, 1.7],
        aligns=["l", "l", "r", "c"],
    )
    note(doc, "※ 소요예산은 해당 월 집행 예정액 기준이며, 확정 금액은 기안 시점에 변동될 수 있습니다.")

    # 2. 주차별 추진 일정
    h1(doc, "2. 주차별 추진 일정")
    caption(doc, "[표 2] 10월 주차별 주요 일정")
    make_table(
        doc,
        ["주차", "기간", "주요 일정", "담당"],
        [
            ("1주차", "10. 1.(목)~10. 4.(일)",
             "북큐레이션 전시 교체 / 희망도서 1주차 접수 마감(10. 4.) / 9월 실적 통계 집계 착수",
             "DM-01·DM-03"),
            ("2주차", "10. 5.(월)~10. 11.(일)",
             "9월 실적 결과보고 기안(10. 7.) / 독서진흥행사 ①②회차 운영(10. 10.) / "
             "11월 정기수서 후보 선정 착수",
             "기획담당·DM-03·DM-01"),
            ("3주차", "10. 12.(월)~10. 18.(일)",
             "운영위원회 2차 회의(10. 15.) / 공모사업 신청 서류 완료(10. 20. 마감 대비) / "
             "순회문고 3회차 배치(10. 16.)",
             "DM-05·DM-03"),
            ("4주차", "10. 19.(월)~10. 25.(일)",
             "독서진흥행사 ③④회차 운영(10. 24.) / 평생학습 중간 만족도 조사 / "
             "복본조사 결과 확정",
             "DM-03·DM-04·DM-01"),
            ("5주차", "10. 26.(월)~10. 31.(토)",
             "문화가 있는 날 운영(10. 28.) / 10월 강사비 지급 기안(10. 29.) / "
             "장서현황 월간 보고(10. 30.) / 11월 계획 수립",
             "DM-03·DM-04·DM-01·기획담당"),
        ],
        widths=[1.5, 3.3, 9.2, 3.0],
        aligns=["c", "c", "l", "c"],
    )

    doc.add_page_break()

    # 3. 도메인별 상세
    h1(doc, "3. 도메인별 상세 계획", space_before=0)

    # DM-01
    h2(doc, "가. DM-01 장서개발")
    make_table(
        doc,
        ["구분", "내용"],
        [
            ("추진 목표", "희망도서 주간 처리 정상 운영, 11월 정기수서 후보 400종 확보"),
            ("담당 리프", "B-01 수서 / B-02 희망도서 / B-03 복본조사 / B-04-H 자료조직 / B-05 장서균형"),
            ("소요예산", "금3,600,000원 (종합자료실운영 · 도서구입비)"),
            ("산출물", "희망도서 구입 건의 기안 4건, 정기수서 후보목록 1건, 장서현황 월간보고 1건"),
        ],
        widths=[2.6, 14.4],
        aligns=["c", "l"],
    )
    caption(doc, "[표 3] DM-01 세부 업무")
    make_table(
        doc,
        ["일자", "업무", "세부 내용", "산출물"],
        [
            ("매주 화", "희망도서 심의", "주간 접수분 R-01~R-10 기준 적용, 복본조사 의뢰", "구입 건의 기안"),
            ("10. 5.~10. 16.", "정기수서 후보 선정", "SEOJI 신간 카탈로그 기반 KDC 배분표 적용, 400종 도출", "후보목록 xlsx"),
            ("10. 19.~10. 23.", "복본조사", "후보 400종 ISBN·서명저자 대조, 상세조사 대상 분리", "복본조사 결과표"),
            ("10. 30.(금)", "장서현황 보고", "10월 말 기준 장서 수·KDC 구성비·증감 집계", "장서현황 보고 기안"),
        ],
        widths=[2.6, 3.2, 8.2, 3.0],
        aligns=["c", "l", "l", "l"],
    )

    # DM-03
    h2(doc, "나. DM-03 독서문화")
    make_table(
        doc,
        ["구분", "내용"],
        [
            ("월간 테마", "「가을, 한 문장에 머물다 — 시(詩)와 함께 걷는 도서관」"),
            ("담당 리프", "D-01 독서동아리 / D-02 행사기획 / D-03 강사섭외 / D-06 순회문고"),
            ("소요예산", "금1,430,000원 (독서문화프로그램운영 · 일반수용비 830,000 + 강사수당 600,000)"),
            ("산출물", "행사 운영계획 1건, 홍보물 3종, 운영일지 6건, 결과보고 1건"),
        ],
        widths=[2.6, 14.4],
        aligns=["c", "l"],
    )
    caption(doc, "[표 4] 10월 독서진흥행사 운영 계획")
    make_table(
        doc,
        ["순", "행사명", "일시", "장소", "대상·규모", "예산(원)"],
        [
            ("1", "가을 시(詩) 북큐레이션", "10. 1.(목)~10. 31.(토) 상시", "종합자료실", "전 연령 자유관람", "0"),
            ("2", "어린이 시 그림책 전시", "10. 1.(목)~10. 31.(토) 상시", "어린이자료실", "유아·초등 자유관람", "0"),
            ("3", "시 필사 워크숍", "10. 10.(토) 14:00~16:00", "강의실 1", "성인 20명", "230,000"),
            ("4", "나만의 시집 만들기", "10. 10.(토) 10:30~12:00", "강의실 2", "초등 3~6학년 10명", "180,000"),
            ("5", "가을밤 시 낭송회", "10. 24.(토) 19:00~20:30", "1층 로비", "전 연령 40명", "320,000"),
            ("6", "문화가 있는 날 — 시노래 공연", "10. 28.(수) 15:00~16:00", "1층 로비", "전 연령 60명", "700,000"),
            ("합  계", "", "", "", "", "1,430,000"),
        ],
        widths=[1.1, 4.4, 4.2, 2.1, 3.0, 2.2],
        aligns=["c", "l", "c", "c", "c", "r"],
        emphasize_last_row=True,
    )
    note(doc, "※ 6번 시노래 공연은 섭외비 견적이 계상액을 초과하여 5절 확인·승인 항목 2번으로 상신하였습니다.")

    # DM-04
    h2(doc, "다. DM-04 평생학습")
    caption(doc, "[표 5] 하반기 강좌 10월 운영 현황")
    make_table(
        doc,
        ["강좌명", "요일·시간", "회차", "정원/등록", "강사", "10월 강사비(원)"],
        [
            ("한 줄 인문학", "화 10:00~12:00", "6~9회차", "20 / 18", "김○○", "600,000"),
            ("어린이 그림책 놀이터", "목 16:00~18:00", "6~9회차", "10 / 10", "박○○", "600,000"),
            ("시니어 스마트폰 활용", "금 14:00~16:00", "6~9회차", "15 / 13", "이○○", "600,000"),
            ("합  계", "", "", "45 / 41", "", "1,800,000"),
        ],
        widths=[4.0, 3.2, 2.0, 2.2, 2.2, 3.4],
        aligns=["l", "c", "c", "c", "c", "r"],
        emphasize_last_row=True,
    )
    body(doc, "가. 10월 4주차에 중간 만족도 조사를 실시하고 결과를 C-06으로 이관합니다.")
    body(doc, "나. 「시니어 스마트폰 활용」 등록률 87%로 추가 모집은 실시하지 않습니다.")

    # DM-02 / DM-05
    h2(doc, "라. DM-02 이용자응대 · DM-05 홍보협력")
    make_table(
        doc,
        ["도메인", "업무", "일정", "비고"],
        [
            ("DM-02", "키오스크 이용안내·추천 응대", "상시", "월간 응대 건수 집계 후 보고"),
            ("DM-02", "상호대차 접수·발송", "상시(일 평균 12건)", "지혜의방 반납 예외 안내 유지"),
            ("DM-02", "하반기 이용자 만족도 조사", "10. 5.~10. 30.", "표본 100명 목표"),
            ("DM-05", "도서관 소식지 4분기호 편집", "10. 12.~10. 30.", "11월 초 발행"),
            ("DM-05", "운영위원회 2차 회의", "10. 15.(목) 15:00", "안건자료·회의록 A-01 작성"),
            ("DM-05", "공모사업 신청", "10. 20.(화) 마감", "신청 여부 승인 필요(5절 3번)"),
        ],
        widths=[2.0, 5.4, 4.0, 5.6],
        aligns=["c", "l", "c", "l"],
    )

    # 4. 예산
    doc.add_page_break()
    h1(doc, "4. 예산 소요 개요", space_before=0)
    caption(doc, "[표 6] 예산과목별 10월 소요액 및 잔액 (2026. 9. 25. 기준)")
    make_table(
        doc,
        ["단위과제카드", "사업항목", "예산목", "예산액(원)", "기집행(원)", "10월 소요(원)", "잔액(원)"],
        BUDGET_ROWS,
        widths=[3.0, 3.4, 3.0, 2.0, 2.0, 2.0, 2.0],
        aligns=["l", "l", "l", "r", "r", "r", "r"],
        font_size=8.5,
        head_size=8.5,
        emphasize_last_row=True,
    )
    body(doc, "가. 10월 소요액 합계는 금7,230,000원이며, 집행 후 잔액은 금16,149,040원입니다.")
    body(doc, "나. 도서구입비 잔액이 11월 정기수서 소요 예상액에 미달하여 유용 검토가 필요합니다"
              "(5절 1번).")
    body(doc, "다. 모든 집행은 기안 결재 후 A-03 예산관리 에이전트에 결의 기록을 반영합니다.")

    # 5. 확인·승인
    h1(doc, "5. 최고관리자 확인 · 승인 필요 항목")
    caption(doc, "[표 7] 확인·승인 요청 사항 (3건)")
    make_table(
        doc,
        ["연번", "출처", "항목", "내용", "회신 기한"],
        ESCALATIONS,
        widths=[1.2, 2.6, 3.8, 6.8, 2.6],
        aligns=["c", "c", "l", "l", "c"],
        font_size=8.5,
    )

    footer_note(doc, "본 문서는 초안이며 최고관리자 검토·확정 전까지 확정본이 아닙니다. "
                     "ㅣ 창녕도서관 기획업무팀 기획담당 ㅣ 2026. 9. 25.")

    path = os.path.join(OUT_DIR, "[양식샘플]2026년_10월_도서관_월간업무계획.docx")
    doc.save(path)
    return path


# ══════════════════════════════════════════════════════════════
# 문서 B — 최고관리자용 요약본 (1~2쪽)
# ══════════════════════════════════════════════════════════════
def build_summary():
    doc = new_doc()

    doc_title(doc, "2026년 10월 월간 업무계획 요약보고",
              "창녕도서관 기획업무팀 기획담당 ㅣ 대상 기간 2026. 10. 1. ~ 10. 31. ㅣ 작성일 2026. 9. 25.")

    # 한눈에 보기
    h1(doc, "한눈에 보기", space_before=6)
    kpi_row(doc, [
        ("운영 사업", "12", "종"),
        ("10월 소요예산", "723", "만원"),
        ("예산 집행률", "70", "%"),
        ("참여 예정", "630", "명"),
        ("승인 필요", "3", "건"),
    ])
    note(doc, "※ 집행률은 연간 예산액 대비 10월 집행 후 누계 기준(₩37,590,960 / ₩53,740,000).")

    # 이달의 중점
    h1(doc, "이달의 중점 추진 사항")
    body(doc, "1. 「독서의 달」 후속 독서진흥행사 6종 운영 — 가을 시(詩) 테마, 참여 예정 130명")
    body(doc, "2. 11월 정기도서수서 사전 준비 완료 — 후보 400종 선정 및 복본조사 10월 내 종료")
    body(doc, "3. 하반기 평생학습 강좌 3종 중간 점검 — 등록률 91%, 중간 만족도 조사 실시")

    # 도메인 요약
    h1(doc, "도메인별 계획 요약")
    make_table(
        doc,
        ["도메인", "핵심 업무", "소요예산(원)", "확인필요"],
        DOMAIN_SUMMARY,
        widths=[2.8, 9.4, 2.6, 1.7],
        aligns=["l", "l", "r", "c"],
        font_size=8.5,
        head_size=8.5,
    )

    # 예산 요약
    h1(doc, "예산 요약")
    make_table(
        doc,
        ["단위과제카드", "예산액(원)", "기집행(원)", "10월 소요(원)", "잔액(원)", "집행률"],
        [
            ("독서문화프로그램운영", "8,740,000", "4,460,960", "1,830,000", "2,449,040", "72%"),
            ("평생학습프로그램운영", "9,000,000", "4,500,000", "1,800,000", "2,700,000", "70%"),
            ("종합자료실운영", "36,000,000", "21,400,000", "3,600,000", "11,000,000", "69%"),
            ("합  계", "53,740,000", "30,360,960", "7,230,000", "16,149,040", "70%"),
        ],
        widths=[4.4, 2.6, 2.6, 2.6, 2.6, 2.2],
        aligns=["l", "r", "r", "r", "r", "c"],
        font_size=8.5,
        head_size=8.5,
        emphasize_last_row=True,
    )

    # 결정 필요
    h1(doc, "최고관리자 결정 필요 사항")
    make_table(
        doc,
        ["연번", "항목", "결정 요청 내용", "회신 기한"],
        [
            ("1", "정기수서 예산 부족", "부족액 ₩2,200,000 인접 사업항목 유용 승인 여부", "10. 10.(토)"),
            ("2", "공연 섭외비 초과", "초과분 ₩150,000 전용 또는 프로그램 축소 중 택일", "10. 8.(목)"),
            ("3", "공모사업 신청 여부", "사업비 ₩5,000,000 규모 — 신청 여부 결정", "10. 13.(화)"),
        ],
        widths=[1.2, 3.6, 9.6, 2.6],
        aligns=["c", "l", "l", "c"],
        font_size=9,
    )

    body(doc, "※ 각 항목의 산정 근거와 대안 비교는 전체본 5절(최고관리자 확인·승인 필요 항목)에 "
              "정리되어 있습니다.", size=8.5, color=C_MUTED)

    footer_note(doc, "본 요약본은 전체본(총 4쪽)의 발췌이며, 초안으로 확정 전 문서입니다. "
                     "ㅣ 창녕도서관 기획업무팀 기획담당")

    path = os.path.join(OUT_DIR, "[양식샘플]2026년_10월_월간업무계획_요약보고.docx")
    doc.save(path)
    return path


if __name__ == "__main__":
    os.makedirs(OUT_DIR, exist_ok=True)
    print("생성:", build_full())
    print("생성:", build_summary())
