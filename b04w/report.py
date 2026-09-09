"""출력물 생성 — records.mrk / records.json / 처리현황.xlsx.

.mrk 형식은 outputs/20260902_b04w_kormarc_test/records_normalized.txt 를 따른다.
LDR·008의 공백은 '#'로 치환해 적는다 — 말미 공백이 편집기·git에 소리 없이 지워지기 때문이다.
"""
from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from fixed_fields import hashify
from models import Result

LEGEND = "※ LDR·008의 #는 공백 1칸이다(적재 시 복원)"

STATUS_LABEL = {
    "ok": "정상",
    "needs_info": "정보부족",
    "qa_failed": "QA실패",
    "error": "오류",
}


def render_record(result: Result) -> str:
    record = result.record
    if record is None:
        return ""
    lines = [f"LDR    {hashify(record.leader)}"]
    for tag in ("001", "005", "007", "008"):
        value = record.control.get(tag, "")
        if not value:
            continue
        lines.append(f"{tag}    {hashify(value) if tag == '008' else value}")
    lines += [field.render() for field in record.fields]
    return "\n".join(lines)


def write_mrk(results: list[Result], path: Path, *, title: str = "") -> None:
    out = [
        f"# B-04-W 자료조직 산출물 — {title or datetime.now().strftime('%Y-%m-%d')}",
        f"# 생성: {datetime.now():%Y-%m-%d %H:%M:%S} / 총 {len(results)}건",
        f"# {LEGEND}",
        "# ⚠ 사서 검토 전 초안이다. needs_info·QA실패 표시가 붙은 건은 확인이 남아 있다.",
        "",
    ]
    for result in results:
        out.append("=" * 80)
        out.append(f"[{result.book.seq}] {result.book.title}   ({STATUS_LABEL.get(result.status, result.status)})")
        out.append("=" * 80)
        if result.record is None:
            out.append(f"# 레코드 생성 실패: {result.error}")
            out.append("")
            continue
        out.append(render_record(result))
        if result.author_mark_derivation:
            out.append(f"# 저자기호 산출: {result.author_mark_derivation}")
        if result.classification and result.classification.kdc_path:
            out.append(f"# 분류 전개: {result.classification.kdc_path}")
        for note in result.notes:
            out.append(f"# 메모: {note}")
        for item in result.needs_info:
            out.append(f"# 확인필요: {item}")
        for failure in result.qa_failures:
            out.append(f"# QA실패: {failure}")
        out.append("")
    path.write_text("\n".join(out), encoding="utf-8")


def write_json(results: list[Result], path: Path) -> None:
    """B-04-H 배치 처리·장서 DB INSERT가 그대로 받아쓸 수 있는 구조."""
    payload = []
    for result in results:
        record = result.record
        payload.append(
            {
                "seq": result.book.seq,
                "status": result.status,
                "isbn": result.book.isbn,
                "reg_nos": result.book.reg_nos,
                "title": result.book.title,
                "author": result.book.author_raw,
                "publisher": result.book.publisher,
                "pub_year": result.book.pub_year,
                "room": result.book.room,
                "loc_mark": {"성인": "", "어린이": "J", "유아": "유"}.get(result.book.room, ""),
                "call_no": result.call_no,
                "vol": result.book.vol,
                "ctrl_no": (record.control.get("001") if record else "") or "",
                "kdc": result.classification.kdc if result.classification else "",
                "kdc_rationale": result.classification.kdc_rationale if result.classification else "",
                "author_mark": result.author_mark,
                "author_mark_derivation": result.author_mark_derivation,
                "marc": {
                    "LDR": record.leader if record else "",
                    "control": record.control if record else {},
                    "fields": [
                        {"tag": f.tag, "ind": f.ind, "value": f.value} for f in (record.fields if record else [])
                    ],
                },
                "qa_failures": result.qa_failures,
                "needs_info": result.needs_info,
                "notes": result.notes,
                "error": result.error,
            }
        )
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_status_xlsx(results: list[Result], path: Path) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "처리현황"
    headers = [
        "순번", "상태", "등록번호", "ISBN", "서명", "저자", "출판사", "발행년",
        "자료실", "KDC", "청구기호", "저자기호 산출", "분류 근거", "확인필요", "QA실패", "메모",
    ]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="DDE5F0")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    fills = {
        "needs_info": PatternFill("solid", fgColor="FFF3CD"),
        "qa_failed": PatternFill("solid", fgColor="F8D7DA"),
        "error": PatternFill("solid", fgColor="F5C6CB"),
    }
    for result in results:
        sheet.append([
            result.book.seq,
            STATUS_LABEL.get(result.status, result.status),
            ", ".join(result.book.reg_nos),
            result.book.isbn,
            result.book.title,
            result.book.author_raw,
            result.book.publisher,
            result.book.pub_year,
            result.book.room,
            result.classification.kdc if result.classification else "",
            result.call_no,
            result.author_mark_derivation,
            result.classification.kdc_rationale if result.classification else "",
            "\n".join(result.needs_info),
            "\n".join(result.qa_failures),
            "\n".join(result.notes) + (f"\n{result.error}" if result.error else ""),
        ])
        fill = fills.get(result.status)
        if fill:
            for cell in sheet[sheet.max_row]:
                cell.fill = fill

    widths = [6, 10, 18, 16, 32, 18, 14, 8, 8, 10, 18, 34, 46, 40, 40, 40]
    for pos, width in enumerate(widths, start=1):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(pos)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    book.save(path)


def summary_line(results: list[Result]) -> str:
    counts: dict[str, int] = {}
    for result in results:
        counts[result.status] = counts.get(result.status, 0) + 1
    parts = [f"{STATUS_LABEL.get(k, k)} {v}건" for k, v in sorted(counts.items())]
    return f"총 {len(results)}건 — " + " / ".join(parts)
