"""처리현황.csv → 처리현황.xlsx. 사서가 실제로 여는 형식으로 한 부 더 남긴다."""
from __future__ import annotations

import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATUS_FILL = {
    "ok": "E8F5E9",
    "needs_info": "FFF8E1",
    "qa_failed": "FFEBEE",
    "error": "FFCDD2",
}
WIDTHS = [5, 12, 14, 40, 26, 16, 10, 16, 10, 18, 11, 8, 40, 40, 46]


def main(source: str) -> int:
    path = Path(source)
    with open(path, encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "처리현황"
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="37474F")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    status_col = rows[0].index("상태") + 1
    for pos in range(2, sheet.max_row + 1):
        status = sheet.cell(row=pos, column=status_col).value
        color = STATUS_FILL.get(status)
        if color:
            sheet.cell(row=pos, column=status_col).fill = PatternFill("solid", fgColor=color)
        for cell in sheet[pos]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for pos, width in enumerate(WIDTHS[: sheet.max_column], start=1):
        sheet.column_dimensions[get_column_letter(pos)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    out = path.with_suffix(".xlsx")
    book.save(out)
    print(f"{sheet.max_row - 1}행 → {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
